import asyncio
from io import BytesIO
import unittest
from docx import Document
from openpyxl import Workbook
from starlette.datastructures import UploadFile
from app.models import ExportRequest, SessionSchedule, Assignment, TeacherDutyCount
from app.exporter import generate_docx
from app.main import import_list


class ExportTests(unittest.TestCase):
    def test_xlsx_import_returns_first_sheet_rows(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Room Number", "Capacity"])
        worksheet.append(["Room 101", 30])
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)

        result = asyncio.run(import_list(UploadFile(filename="rooms.xlsx", file=stream)))

        self.assertEqual(result["rows"][0], ["Room Number", "Capacity"])
        self.assertEqual(result["rows"][1], ["Room 101", "30"])

    def test_schedule_table_has_two_teacher_shifts_and_free_list(self):
        data = ExportRequest(
            school_name="Example School",
            exam_name="Term Exams",
            schedule=[
                SessionSchedule(
                    session_id="am",
                    date="2026-10-12",
                    time_slot="09:00 AM - 12:00 PM",
                    session_name="Morning",
                    assignments=[Assignment(room="101", teacher="A"), Assignment(room="102", teacher="B")]
                ),
                SessionSchedule(
                    session_id="pm",
                    date="2026-10-12",
                    time_slot="01:00 PM - 04:00 PM",
                    session_name="Afternoon",
                    assignments=[Assignment(room="101", teacher="C"), Assignment(room="102", teacher="D")]
                )
            ],
            fairness_matrix=[TeacherDutyCount(teacher=name, total_duties=1) for name in "ABCD"]
            + [TeacherDutyCount(teacher="E", total_duties=0)]
        )

        document = Document(generate_docx(data))
        table = document.tables[0]
        headers = [cell.text for cell in table.rows[0].cells]
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertEqual(headers[0], "Room No.")
        self.assertIn("Duty Shift 1 - Teacher Name", headers[1])
        self.assertIn("Duty Shift 2 - Teacher Name", headers[2])
        self.assertIn("Remaining and Free Teachers", paragraph_text)
        self.assertIn("Shift 1 standby (3)", paragraph_text)
        self.assertIn("Shift 2 standby (3)", paragraph_text)
        self.assertIn("Full-day standby (1)", paragraph_text)
        self.assertIn("E", paragraph_text)
        self.assertIn("C; D; E", paragraph_text)
        self.assertNotIn("Fairness Audit Matrix", paragraph_text)
        self.assertEqual(len(document.tables), 1)
        self.assertFalse(any(paragraph.text == "" for paragraph in document.paragraphs))


if __name__ == "__main__":
    unittest.main()

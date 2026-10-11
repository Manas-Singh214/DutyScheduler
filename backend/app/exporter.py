import io
from typing import Dict
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from .models import ExportRequest

def set_cell_background(cell, fill_hex: str):
    """Utility to set cell shading color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Utility to set cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def group_sessions_into_shifts(schedule):
    date_to_sessions: Dict[str, list] = {}
    for session in schedule:
        date_to_sessions.setdefault(session.date, []).append(session)

    has_multi = any(len(s_list) > 1 for s_list in date_to_sessions.values())

    day_groups = []
    if has_multi:
        for date, s_list in date_to_sessions.items():
            for i in range(0, len(s_list), 2):
                s1 = s_list[i]
                s2 = s_list[i + 1] if i + 1 < len(s_list) else None
                day_groups.append({
                    "title": f"Exam Date: {date}",
                    "date": date,
                    "shift1": s1,
                    "shift2": s2
                })
    else:
        if len(schedule) % 2 == 0 and len(schedule) > 1:
            for i in range(0, len(schedule), 2):
                s1 = schedule[i]
                s2 = schedule[i + 1]
                title = f"Exam Dates: {s1.date} & {s2.date}" if s1.date != s2.date else f"Exam Date: {s1.date}"
                day_groups.append({
                    "title": title,
                    "date": s1.date,
                    "shift1": s1,
                    "shift2": s2
                })
        else:
            for i in range(0, len(schedule), 2):
                s1 = schedule[i]
                s2 = schedule[i + 1] if i + 1 < len(schedule) else None
                day_groups.append({
                    "title": f"Exam Date: {s1.date}",
                    "date": s1.date,
                    "shift1": s1,
                    "shift2": s2
                })
    return day_groups

def generate_docx(data: ExportRequest) -> io.BytesIO:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    # Theme Colors
    NAVY_HEX = "1E293B"
    LIGHT_BG = "F8FAFC"

    # Header - School Name
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(data.school_name.upper())
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Subtitle - Exam Name
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(data.exam_name)
    sub_run.font.name = 'Calibri'
    sub_run.font.size = Pt(14)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    # Document Label
    label_p = doc.add_paragraph()
    label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_run = label_p.add_run("OFFICIAL EXAM DUTY INVIGILATION SCHEDULE")
    label_run.font.name = 'Calibri'
    label_run.font.size = Pt(11)
    label_run.font.bold = True
    label_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    all_teachers = [item.teacher for item in data.fairness_matrix]
    day_groups = group_sessions_into_shifts(data.schedule)

    # Generate each exam shift day
    for group in day_groups:
        s1 = group["shift1"]
        s2 = group["shift2"]

        # Date & Shift Heading
        day_p = doc.add_paragraph()
        day_run = day_p.add_run(group['title'])
        day_run.font.name = 'Calibri'
        day_run.font.size = Pt(13)
        day_run.font.bold = True
        day_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        # Sub-heading details
        sub_text = f"Shift 1: {s1.session_name} ({s1.time_slot})"
        if s2:
            sub_text += f"   |   Shift 2: {s2.session_name} ({s2.time_slot})"
        info_p = doc.add_paragraph()
        info_run = info_p.add_run(sub_text)
        info_run.font.name = 'Calibri'
        info_run.font.size = Pt(9.5)
        info_run.font.italic = True
        info_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        # Collect rooms and assignment mappings
        room_list = []
        s1_map = {a.room: a.teacher for a in s1.assignments}
        s2_map = {a.room: a.teacher for a in s2.assignments} if s2 else {}

        for a in s1.assignments:
            if a.room not in room_list:
                room_list.append(a.room)
        if s2:
            for a in s2.assignments:
                if a.room not in room_list:
                    room_list.append(a.room)

        # Combined Table: Room No. | Duty Shift 1 | Duty Shift 2
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        col_widths = [Inches(1.8), Inches(2.55), Inches(2.55)]

        # Header Row
        hdr_cells = table.rows[0].cells
        hdr_s1_label = f"Duty Shift 1 - Teacher Name\n{s1.session_name}"
        hdr_s2_label = f"Duty Shift 2 - Teacher Name\n{s2.session_name}" if s2 else "Duty Shift 2 - Teacher Name\n(No Exam)"
        hdr_titles = ["Room No.", hdr_s1_label, hdr_s2_label]

        for i, title in enumerate(hdr_titles):
            hdr_cells[i].width = col_widths[i]
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(title)
            r.font.name = 'Calibri'
            r.font.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_cell_background(hdr_cells[i], NAVY_HEX)
            set_cell_margins(hdr_cells[i], top=120, bottom=120)

        # Populate Room Rows
        for r_idx, room in enumerate(room_list):
            row_cells = table.add_row().cells
            for i in range(3):
                row_cells[i].width = col_widths[i]

            # Room No.
            p0 = row_cells[0].paragraphs[0]
            r0 = p0.add_run(room)
            r0.font.name = 'Calibri'
            r0.font.bold = True
            r0.font.size = Pt(10)

            # Shift 1 Teacher
            p1 = row_cells[1].paragraphs[0]
            t1 = s1_map.get(room)
            if t1:
                r1_name = p1.add_run(t1 + "\n")
                r1_name.font.name = 'Calibri'
                r1_name.font.bold = True
                r1_name.font.size = Pt(10)
                r1_sig = p1.add_run("Sig: ____________________")
                r1_sig.font.name = 'Calibri'
                r1_sig.font.size = Pt(8.5)
                r1_sig.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
            else:
                r1 = p1.add_run("—")
                r1.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

            # Shift 2 Teacher
            p2 = row_cells[2].paragraphs[0]
            if s2:
                t2 = s2_map.get(room)
                if t2:
                    r2_name = p2.add_run(t2 + "\n")
                    r2_name.font.name = 'Calibri'
                    r2_name.font.bold = True
                    r2_name.font.size = Pt(10)
                    r2_sig = p2.add_run("Sig: ____________________")
                    r2_sig.font.name = 'Calibri'
                    r2_sig.font.size = Pt(8.5)
                    r2_sig.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
                else:
                    r2 = p2.add_run("—")
                    r2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
            else:
                r2 = p2.add_run("— (No Exam)")
                r2.font.italic = True
                r2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

            # Alternating row background
            if r_idx % 2 == 1:
                set_cell_background(row_cells[0], LIGHT_BG)
                set_cell_background(row_cells[1], LIGHT_BG)
                set_cell_background(row_cells[2], LIGHT_BG)

            for cell in row_cells:
                set_cell_margins(cell, top=100, bottom=100)

        # Spacing
        # Calculate Free / Remaining Teachers
        s1_assigned = set(s1_map.values())
        s2_assigned = set(s2_map.values()) if s2 else set()
        duty_all = s1_assigned | s2_assigned

        s1_free = [t for t in all_teachers if t not in s1_assigned]
        s2_free = [t for t in all_teachers if t not in s2_assigned] if s2 else []
        full_free = [t for t in all_teachers if t not in duty_all]

        # Remaining & Free Teachers Section
        free_hdr = doc.add_paragraph()
        free_hdr_run = free_hdr.add_run("Remaining and Free Teachers (Relievers / Standby Duty)")
        free_hdr_run.font.name = 'Calibri'
        free_hdr_run.font.bold = True
        free_hdr_run.font.size = Pt(11)
        free_hdr_run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        free_hdr.paragraph_format.keep_with_next = True
        free_hdr.paragraph_format.space_before = Pt(2)
        free_hdr.paragraph_format.space_after = Pt(0)
        free_values = [f"Shift 1 standby ({len(s1_free)}): {'; '.join(s1_free) or 'None'}"]
        if s2:
            free_values.append(f"Shift 2 standby ({len(s2_free)}): {'; '.join(s2_free) or 'None'}")
        if full_free:
            free_values.append(f"Full-day standby ({len(full_free)}): {'; '.join(full_free)}")
        for line in free_values:
            paragraph = doc.add_paragraph(line)
            paragraph.paragraph_format.space_after = Pt(0)

    # Save to Stream
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

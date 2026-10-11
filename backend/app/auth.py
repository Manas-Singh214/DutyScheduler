import os
import hashlib
import io
import secrets
import smtplib
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from typing import Annotated

import jwt
from fastapi import APIRouter, Depends, File, Header, HTTPException, Request, UploadFile, status
from pydantic import BaseModel, Field
from pwdlib import PasswordHash
from openpyxl import Workbook, load_workbook
from fastapi.responses import StreamingResponse
from sqlalchemy import ForeignKey, String, UniqueConstraint, select
from sqlalchemy.orm import Session, Mapped, mapped_column

from .database import Base, SessionLocal
from .rate_limit import limiter

router = APIRouter(prefix="/auth", tags=["school accounts"])
password_hash = PasswordHash.recommended()
SECRET_KEY = os.getenv("APP_SECRET_KEY", "local-development-key-change-before-deploy")
FRONTEND_URL = os.getenv(
    "FRONTEND_URL",
    os.getenv("RENDER_EXTERNAL_URL", "http://localhost:5173"),
).rstrip("/")
if os.getenv("APP_ENV") == "production":
    if SECRET_KEY == "local-development-key-change-before-deploy":
        raise RuntimeError("Set APP_SECRET_KEY to a unique secret before production.")
    if not os.getenv("DATABASE_URL", "").startswith(("postgres://", "postgresql://", "postgresql+psycopg://")):
        raise RuntimeError("Production requires a persistent PostgreSQL DATABASE_URL.")
    if not os.getenv("SMTP_HOST") or not os.getenv("SMTP_FROM"):
        raise RuntimeError("Production requires SMTP_HOST and SMTP_FROM for account email.")


class School(Base):
    __tablename__ = "schools"
    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(200))


class User(Base):
    __tablename__ = "users"
    id: Mapped[int] = mapped_column(primary_key=True)
    school_id: Mapped[int] = mapped_column(ForeignKey("schools.id", ondelete="CASCADE"), index=True)
    email: Mapped[str] = mapped_column(String(320), unique=True, index=True)
    password_hash: Mapped[str] = mapped_column(String(500))
    role: Mapped[str] = mapped_column(String(20), default="staff")
    verified: Mapped[bool] = mapped_column(default=False)


class Invitation(Base):
    __tablename__ = "invitations"
    id: Mapped[int] = mapped_column(primary_key=True)
    school_id: Mapped[int] = mapped_column(ForeignKey("schools.id", ondelete="CASCADE"), index=True)
    email: Mapped[str] = mapped_column(String(320), index=True)
    token_hash: Mapped[str] = mapped_column(String(64), unique=True)
    expires_at: Mapped[datetime]
    accepted: Mapped[bool] = mapped_column(default=False)


class DutyHistory(Base):
    __tablename__ = "duty_history"
    __table_args__ = (UniqueConstraint("school_id", "teacher_name"),)
    id: Mapped[int] = mapped_column(primary_key=True)
    school_id: Mapped[int] = mapped_column(ForeignKey("schools.id", ondelete="CASCADE"), index=True)
    teacher_name: Mapped[str] = mapped_column(String(200))
    total_duties: Mapped[int] = mapped_column(default=0)


class CompletedExam(Base):
    __tablename__ = "completed_exams"
    __table_args__ = (UniqueConstraint("school_id", "exam_id"),)
    id: Mapped[int] = mapped_column(primary_key=True)
    school_id: Mapped[int] = mapped_column(ForeignKey("schools.id", ondelete="CASCADE"), index=True)
    exam_id: Mapped[str] = mapped_column(String(100))
    exam_name: Mapped[str] = mapped_column(String(200))
    completed_at: Mapped[datetime] = mapped_column(default=lambda: datetime.now(timezone.utc))


class ExamDuty(Base):
    __tablename__ = "exam_duties"
    school_id: Mapped[int] = mapped_column(ForeignKey("schools.id", ondelete="CASCADE"), index=True)
    id: Mapped[int] = mapped_column(primary_key=True)
    exam_id: Mapped[str] = mapped_column(String(100), index=True)
    teacher_name: Mapped[str] = mapped_column(String(200))
    duty_count: Mapped[int] = mapped_column(default=0)


class DutyImport(Base):
    __tablename__ = "duty_imports"
    __table_args__ = (UniqueConstraint("school_id", "content_hash"),)
    id: Mapped[int] = mapped_column(primary_key=True)
    school_id: Mapped[int] = mapped_column(ForeignKey("schools.id", ondelete="CASCADE"), index=True)
    content_hash: Mapped[str] = mapped_column(String(64))


class PasswordResetToken(Base):
    __tablename__ = "password_reset_tokens"
    id: Mapped[int] = mapped_column(primary_key=True)
    email: Mapped[str] = mapped_column(String(320), index=True)
    token_hash: Mapped[str] = mapped_column(String(64), unique=True)
    expires_at: Mapped[datetime]
    consumed: Mapped[bool] = mapped_column(default=False)


class Registration(BaseModel):
    school_name: str = Field(min_length=2, max_length=200)
    email: str = Field(min_length=5, max_length=320)
    password: str = Field(min_length=12, max_length=256)


class Credentials(BaseModel):
    email: str
    password: str = Field(min_length=1, max_length=256)


class TokenInput(BaseModel):
    token: str


class PasswordResetInput(TokenInput):
    password: str = Field(min_length=12, max_length=256)


class PasswordResetRequest(BaseModel):
    email: str


class InviteRequest(BaseModel):
    email: str


class HistoryRecord(BaseModel):
    exam_id: str = Field(min_length=1, max_length=100)
    exam_name: str = Field(min_length=1, max_length=200)
    teacher_counts: dict[str, int]


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def normalized_email(value: str) -> str:
    email = value.strip().lower()
    if email.count("@") != 1 or "." not in email.rsplit("@", 1)[-1]:
        raise HTTPException(status_code=422, detail="Enter a valid email address.")
    return email


def create_token(subject: str, purpose: str, expires: timedelta) -> str:
    now = datetime.now(timezone.utc)
    return jwt.encode(
        {"sub": subject, "purpose": purpose, "iat": now, "exp": now + expires},
        SECRET_KEY,
        algorithm="HS256",
    )


def decode_token(token: str, purpose: str) -> str:
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=["HS256"])
    except jwt.PyJWTError as error:
        raise HTTPException(status_code=400, detail="This link is invalid or expired.") from error
    if payload.get("purpose") != purpose or not payload.get("sub"):
        raise HTTPException(status_code=400, detail="This link is invalid or expired.")
    return str(payload["sub"])


def send_email(recipient: str, subject: str, body: str) -> None:
    host = os.getenv("SMTP_HOST")
    sender = os.getenv("SMTP_FROM")
    if not host or not sender:
        raise RuntimeError("Configure SMTP_HOST and SMTP_FROM to enable school account emails.")
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = sender
    message["To"] = recipient
    message.set_content(body)
    with smtplib.SMTP(host, int(os.getenv("SMTP_PORT", "587")), timeout=15) as server:
        server.starttls()
        username = os.getenv("SMTP_USER")
        if username:
            server.login(username, os.getenv("SMTP_PASSWORD", ""))
        server.send_message(message)


def deliver(recipient: str, subject: str, body: str) -> None:
    try:
        send_email(recipient, subject, body)
    except (OSError, RuntimeError, smtplib.SMTPException, ValueError) as error:
        raise HTTPException(status_code=503, detail=f"Email delivery is unavailable: {error}") from error


def get_current_user(
    authorization: Annotated[str | None, Header()] = None,
    db: Session = Depends(get_db),
) -> User:
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Sign in to access this school workspace.")
    user_id = decode_token(authorization[7:], "access")
    user = db.get(User, int(user_id)) if user_id.isdigit() else None
    if user is None or not user.verified:
        raise HTTPException(status_code=401, detail="Your account is not available. Sign in again.")
    return user


def get_admin(user: User = Depends(get_current_user)) -> User:
    if user.role != "admin":
        raise HTTPException(status_code=403, detail="Only a school admin can invite staff.")
    return user


@router.post("/register", status_code=202)
@limiter.limit("5/hour")
def register(request: Request, payload: Registration, db: Session = Depends(get_db)):
    email = normalized_email(payload.email)
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(status_code=409, detail="This email already has an account.")
    school = School(name=payload.school_name.strip())
    db.add(school)
    db.flush()
    user = User(
        school_id=school.id,
        email=email,
        password_hash=password_hash.hash(payload.password),
        role="admin",
        verified=False,
    )
    db.add(user)
    token = create_token(email, "verify-email", timedelta(hours=24))
    deliver(email, "Verify your Exam Duty Allocator account", f"Verify your account: {FRONTEND_URL}/?verify_token={token}")
    db.commit()
    return {"message": "Check your email to verify your school account."}


@router.post("/verify")
@limiter.limit("10/minute")
def verify_email(request: Request, payload: TokenInput, db: Session = Depends(get_db)):
    email = decode_token(payload.token, "verify-email")
    user = db.scalar(select(User).where(User.email == email))
    if user is None:
        raise HTTPException(status_code=400, detail="This verification link is invalid.")
    user.verified = True
    db.commit()
    return {"message": "Email verified. You can now sign in."}


@router.post("/login")
@limiter.limit("10/minute")
def login(request: Request, payload: Credentials, db: Session = Depends(get_db)):
    email = normalized_email(payload.email)
    user = db.scalar(select(User).where(User.email == email))
    if user is None or not password_hash.verify(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Email or password is incorrect.")
    if not user.verified:
        raise HTTPException(status_code=403, detail="Verify your email before signing in.")
    return {
        "access_token": create_token(str(user.id), "access", timedelta(hours=12)),
        "token_type": "bearer",
        "school_name": db.get(School, user.school_id).name,
        "role": user.role,
        "email": user.email,
    }


@router.post("/invite", status_code=202)
@limiter.limit("10/hour")
def invite_staff(request: Request, payload: InviteRequest, admin: User = Depends(get_admin), db: Session = Depends(get_db)):
    email = normalized_email(payload.email)
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(status_code=409, detail="This email already has an account.")
    raw_token = secrets.token_urlsafe(32)
    db.add(Invitation(
        school_id=admin.school_id,
        email=email,
        token_hash=hashlib.sha256(raw_token.encode()).hexdigest(),
        expires_at=datetime.now(timezone.utc) + timedelta(days=3),
    ))
    deliver(email, "Invitation to your school workspace", f"Accept the invitation: {FRONTEND_URL}/?invite_token={raw_token}")
    db.commit()
    return {"message": "Invitation sent."}


@router.post("/invite/accept")
@limiter.limit("10/hour")
def accept_invitation(request: Request, payload: PasswordResetInput, db: Session = Depends(get_db)):
    token_hash = hashlib.sha256(payload.token.encode()).hexdigest()
    invitation = db.scalar(select(Invitation).where(Invitation.token_hash == token_hash))
    now = datetime.now(timezone.utc)
    if invitation is None or invitation.accepted or invitation.expires_at.replace(tzinfo=timezone.utc) < now:
        raise HTTPException(status_code=400, detail="This invitation is invalid or expired.")
    if db.scalar(select(User).where(User.email == invitation.email)):
        raise HTTPException(status_code=409, detail="This email already has an account.")
    db.add(User(
        school_id=invitation.school_id,
        email=invitation.email,
        password_hash=password_hash.hash(payload.password),
        role="staff",
        verified=True,
    ))
    invitation.accepted = True
    db.commit()
    return {"message": "Account created. You can now sign in."}


@router.post("/password-reset", status_code=202)
@limiter.limit("5/hour")
def request_password_reset(request: Request, payload: PasswordResetRequest, db: Session = Depends(get_db)):
    email = normalized_email(payload.email)
    user = db.scalar(select(User).where(User.email == email, User.verified.is_(True)))
    if user:
        token = secrets.token_urlsafe(32)
        db.query(PasswordResetToken).filter(
            PasswordResetToken.email == email,
        ).delete(synchronize_session=False)
        db.add(PasswordResetToken(
            email=email,
            token_hash=hashlib.sha256(token.encode()).hexdigest(),
            expires_at=datetime.now(timezone.utc) + timedelta(minutes=30),
        ))
        deliver(email, "Reset your Exam Duty Allocator password", f"Reset your password: {FRONTEND_URL}/?reset_token={token}")
        db.commit()
    return {"message": "If that verified account exists, a reset link has been sent."}


@router.post("/password-reset/confirm")
@limiter.limit("10/hour")
def confirm_password_reset(request: Request, payload: PasswordResetInput, db: Session = Depends(get_db)):
    token_hash = hashlib.sha256(payload.token.encode()).hexdigest()
    reset_token = db.scalar(select(PasswordResetToken).where(PasswordResetToken.token_hash == token_hash))
    now = datetime.now(timezone.utc)
    if (
        reset_token is None
        or reset_token.consumed
        or reset_token.expires_at.replace(tzinfo=timezone.utc) < now
    ):
        raise HTTPException(status_code=400, detail="This reset link is invalid or expired.")
    user = db.scalar(select(User).where(User.email == reset_token.email))
    if user is None or not user.verified:
        raise HTTPException(status_code=400, detail="This reset link is invalid.")
    user.password_hash = password_hash.hash(payload.password)
    reset_token.consumed = True
    db.commit()
    return {"message": "Password updated. Sign in with your new password."}


@router.get("/me")
def account_info(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return {
        "email": user.email,
        "role": user.role,
        "school_name": db.get(School, user.school_id).name,
    }


@router.get("/history")
def duty_history(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.scalars(select(DutyHistory).where(DutyHistory.school_id == user.school_id)).all()
    return {row.teacher_name: row.total_duties for row in rows}


@router.post("/history/record")
def record_completed_exam(
    payload: HistoryRecord,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if any(count < 0 for count in payload.teacher_counts.values()):
        raise HTTPException(status_code=422, detail="Teacher duty counts cannot be negative.")
    if db.scalar(select(CompletedExam).where(
        CompletedExam.school_id == user.school_id,
        CompletedExam.exam_id == payload.exam_id,
    )):
        raise HTTPException(status_code=409, detail="This exam has already been added to history.")
    db.add(CompletedExam(school_id=user.school_id, exam_id=payload.exam_id, exam_name=payload.exam_name))
    for name, count in payload.teacher_counts.items():
        if name.strip() and count > 0:
            db.add(ExamDuty(
                school_id=user.school_id,
                exam_id=payload.exam_id,
                teacher_name=name.strip(),
                duty_count=count,
            ))
    for name, count in payload.teacher_counts.items():
        teacher_name = name.strip()
        if not teacher_name:
            continue
        row = db.scalar(select(DutyHistory).where(
            DutyHistory.school_id == user.school_id,
            DutyHistory.teacher_name == teacher_name,
        ))
        if row:
            row.total_duties += count
        else:
            db.add(DutyHistory(
                school_id=user.school_id,
                teacher_name=teacher_name,
                total_duties=count,
            ))
    db.commit()
    return {"message": "Exam duties added to this school's history."}


@router.post("/history/import")
@limiter.limit("5/hour")
async def import_duty_history(
    request: Request,
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload an .xlsx or .xlsm history workbook.")
    contents = await file.read(10 * 1024 * 1024 + 1)
    if len(contents) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Workbook exceeds the 10 MB limit.")
    content_hash = hashlib.sha256(contents).hexdigest()
    if db.scalar(select(DutyImport).where(
        DutyImport.school_id == user.school_id,
        DutyImport.content_hash == content_hash,
    )):
        raise HTTPException(status_code=409, detail="This workbook has already been imported.")
    try:
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"Could not read workbook: {error}") from error
    if not rows:
        raise HTTPException(status_code=400, detail="The workbook is empty.")
    headers = [str(value or "").strip().lower() for value in rows[0]]
    teacher_col = next((i for i, value in enumerate(headers) if value in {"teacher", "teacher name", "staff", "staff name"}), None)
    count_col = next((i for i, value in enumerate(headers) if value in {"total duties", "duties", "duty count", "total shifts"}), None)
    if teacher_col is None or count_col is None:
        raise HTTPException(status_code=400, detail="The first sheet needs 'Teacher Name' and 'Total Duties' columns.")
    imported: dict[str, int] = {}
    for row in rows[1:]:
        if teacher_col >= len(row) or count_col >= len(row) or not row[teacher_col]:
            continue
        teacher = str(row[teacher_col]).strip()
        try:
            count = int(row[count_col])
        except (TypeError, ValueError) as error:
            raise HTTPException(status_code=400, detail=f"Invalid duty total for {teacher}.") from error
        if count < 0:
            raise HTTPException(status_code=400, detail=f"Duty total cannot be negative for {teacher}.")
        imported[teacher] = imported.get(teacher, 0) + count
    if not imported:
        raise HTTPException(status_code=400, detail="No teacher duty totals were found.")
    for teacher, count in imported.items():
        row = db.scalar(select(DutyHistory).where(
            DutyHistory.school_id == user.school_id,
            DutyHistory.teacher_name == teacher,
        ))
        if row:
            row.total_duties += count
        else:
            db.add(DutyHistory(school_id=user.school_id, teacher_name=teacher, total_duties=count))
    db.add(DutyImport(school_id=user.school_id, content_hash=content_hash))
    db.commit()
    return {"message": f"Added past-duty totals for {len(imported)} teachers."}


@router.get("/history/export.xlsx")
def export_duty_history(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    totals = db.scalars(
        select(DutyHistory).where(DutyHistory.school_id == user.school_id).order_by(DutyHistory.teacher_name)
    ).all()
    exams = db.scalars(
        select(CompletedExam).where(CompletedExam.school_id == user.school_id).order_by(CompletedExam.completed_at)
    ).all()
    workbook = Workbook()
    totals_sheet = workbook.active
    totals_sheet.title = "Teacher Totals"
    totals_sheet.append(["Teacher Name", "Total Duties"])
    for row in totals:
        totals_sheet.append([row.teacher_name, row.total_duties])
    log_sheet = workbook.create_sheet("Exam Log")
    log_sheet.append(["Exam", "Teacher Name", "Duties"])
    duties = {exam.exam_id: [] for exam in exams}
    exam_ids = list(duties)
    if exam_ids:
        for duty in db.scalars(select(ExamDuty).where(
            ExamDuty.school_id == user.school_id,
            ExamDuty.exam_id.in_(exam_ids),
        )).all():
            duties[duty.exam_id].append(duty)
    for exam in exams:
        for duty in duties[exam.exam_id]:
            log_sheet.append([exam.exam_name, duty.teacher_name, duty.duty_count])
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="school-duty-history.xlsx"'},
    )
